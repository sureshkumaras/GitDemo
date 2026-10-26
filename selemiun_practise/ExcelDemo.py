import openpyxl

book = openpyxl.load_workbook("C:\\Z_Python\\PythonDemo.xlsx")
sheet = book.active
cell = sheet.cell(row=1, column=1)
print(cell.value)

Dict = {}

for i in range(1, sheet.max_row+1):
    if sheet.cell(row=i, column=1).value == "Testcase2" :
        for j in range(2, sheet.max_column+1) :
                Dict[str(sheet.cell(row=1, column=j).value)] = sheet.cell(row=i, column=j).value
            #print(sheet.cell(row=i, column=j).value)
        break


print(Dict)
