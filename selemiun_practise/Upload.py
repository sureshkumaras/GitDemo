import time

import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait


####################################################################################
def update_excel_data(filePath, searchTerm, colName, new_value):
    book = openpyxl.load_workbook(r"C:\\Users\\sures\\Downloads\\download.xlsx")
    sheet = book.active
    dict = {}
    for i in range(1,sheet.max_column+1):
        if "price" == sheet.cell(row=1, column=i).value :
            dict["column"] = i
            break

    for i in range(1, sheet.max_row+1):
        for j in range(1, sheet.max_column+1):
            if "Apple" == sheet.cell(row=i, column=j).value :
                dict["row"] = i
                break
    sheet.cell(row=dict["row"], column=dict["column"]).value = newValue
    book.save(r"C:\\Users\\sures\\Downloads\\download.xlsx")
##########################################################################################

driver = webdriver.Chrome()
driver.maximize_window()
driver.implicitly_wait(2)
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")

file_path = r"C:\Users\sures\Downloads\download.xlsx"
fruit_name = "Apple"
newValue = "990"

#download
driver.find_element(By.ID, "downloadButton").click()
time.sleep(2)

#edit the excel with updated value
update_excel_data(file_path, fruit_name, "price", newValue)

#upload
file_input = driver.find_element(By.XPATH, "//input[@type='file']")
file_input.send_keys(file_path)

#explisit wait
wait = WebDriverWait(driver,5)
toast_locator = (By.CSS_SELECTOR,".Toastify__toast-body div:nth-child(2)")

#wait for the toast message to visible
text = wait.until(expected_conditions.visibility_of_element_located(toast_locator)).text
print(text)

priceColumn = driver.find_element(By.XPATH,"//div[text()='Price']").get_attribute("data-column-id")
actual_price = driver.find_element(By.XPATH,"//div[text()='"+fruit_name+"']/parent::div/parent::div/div[@id='cell-"+priceColumn+"-undefined']").text

assert actual_price == newValue

time.sleep(20)